"""
Independent verifier for Module 4 (Delhi-NCR, MCD consolidated).

Re-derives the expected output straight from the source file, without
importing reconcile_ncr.py, and checks the generated Excel AND PDF.
"""
import sys
import re
import openpyxl

SOURCE, XLSX, PDF = sys.argv[1:4]

MCD_ZONES = {"east delhi municipal corporation", "north delhi municipal corporation",
            "south delhi municipal corporation"}


def norm(t):
    if t is None:
        return ""
    s = re.sub(r"\s+", " ", str(t).replace("\u00a0", " ").replace("\n", " ")).strip().casefold()
    return re.sub(r"[.\-_,'\"]+", "", s)


# ---- read source ------------------------------------------------------------
sws = openpyxl.load_workbook(SOURCE, data_only=True).worksheets[0]
src_rows = []
for r in range(3, sws.max_row + 1):
    nm = sws.cell(r, 2).value
    if not nm or norm(nm) == "total":
        continue
    src_rows.append((str(nm).replace("\n", " ").strip(),
                     int(sws.cell(r, 3).value or 0), int(sws.cell(r, 4).value or 0)))

zones = [r for r in src_rows if norm(r[0]) in MCD_ZONES]
others = [r for r in src_rows if norm(r[0]) not in MCD_ZONES]
assert len(zones) == 3, f"expected 3 MCD zones in source, found {len(zones)}"

mcd_recv = sum(z[1] for z in zones)
mcd_res = sum(z[2] for z in zones)
expected_order = others[:11] + [(("Municipal Corporation of Delhi"), mcd_recv, mcd_res)] + others[11:]
expected_total_recv = sum(r[1] for r in src_rows)
expected_total_res = sum(r[2] for r in src_rows)

fails, checked = [], 0

# ---- check Excel -------------------------------------------------------------
xws = openpyxl.load_workbook(XLSX, data_only=True).worksheets[0]

checked += 1
if xws.cell(1, 1).value is None or "in Delhi-NCR" not in xws.cell(1, 1).value:
    fails.append(f"XLSX A1 missing 'in Delhi-NCR' suffix: {xws.cell(1,1).value!r}")

checked += 1
if xws.cell(14, 2).value != "Municipal Corporation of Delhi":
    fails.append(f"XLSX row 14 (Sr.No 12) should be MCD combined, found "
                 f"{xws.cell(14,2).value!r}")
checked += 2
if xws.cell(14, 1).value != 12:
    fails.append(f"XLSX MCD row S.No expected 12, got {xws.cell(14,1).value!r}")
if (xws.cell(14, 3).value, xws.cell(14, 4).value) != (mcd_recv, mcd_res):
    fails.append(f"XLSX MCD combined values expected ({mcd_recv},{mcd_res}), "
                 f"got ({xws.cell(14,3).value},{xws.cell(14,4).value})")

for i, (name, recv, res) in enumerate(expected_order, 1):
    checked += 3
    r = i + 2
    got_name = xws.cell(r, 2).value
    if norm(got_name) != norm(name):
        fails.append(f"XLSX row {r}: expected '{name}', got '{got_name}'")
    if xws.cell(r, 1).value != i:
        fails.append(f"XLSX row {r}: S.No expected {i}, got {xws.cell(r,1).value}")
    if (xws.cell(r, 3).value, xws.cell(r, 4).value) != (recv, res):
        fails.append(f"XLSX row {r} [{name}]: expected ({recv},{res}), "
                     f"got ({xws.cell(r,3).value},{xws.cell(r,4).value})")

total_row = 3 + len(expected_order)
checked += 2
if xws.cell(total_row, 2).value != "Total":
    fails.append(f"XLSX row {total_row} expected 'Total', got {xws.cell(total_row,2).value!r}")
if (xws.cell(total_row, 3).value, xws.cell(total_row, 4).value) != (expected_total_recv, expected_total_res):
    fails.append(f"XLSX Total expected ({expected_total_recv},{expected_total_res}), "
                 f"got ({xws.cell(total_row,3).value},{xws.cell(total_row,4).value})")

# ---- check PDF (text extraction) --------------------------------------------
try:
    import pdfplumber
    with pdfplumber.open(PDF) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        tables = []
        for p in pdf.pages:
            for t in p.extract_tables():
                tables.extend(t)
except Exception as e:
    fails.append(f"Could not read PDF: {e}")
    text, tables = "", []

checked += 1
if "Municipal Corporation of Delhi" not in text:
    fails.append("PDF text does not contain 'Municipal Corporation of Delhi'")

checked += 1
found_mcd_row = any(row and len(row) > 3 and row[1] and "Municipal Corporation of Delhi" in row[1]
                    and row[0] == "12" for row in tables)
if not found_mcd_row:
    fails.append("PDF table: no row with S.No 12 = 'Municipal Corporation of Delhi' found")

checked += 1
found_total = any(row and len(row) > 3 and row[1] == "Total"
                  and row[2] == str(expected_total_recv) for row in tables)
if not found_total:
    fails.append(f"PDF table: Total row with Received={expected_total_recv} not found")

print(f"assertions checked: {checked}")
if fails:
    print(f"FAILURES: {len(fails)}")
    for f in fails:
        print("   x", f)
    sys.exit(1)
print("ALL CHECKS PASSED")
